import tempfile
import unittest
from pathlib import Path

import docx
import openpyxl

from lgpd_scanner.patterns import find_matches, mask, valid_cnpj, valid_cpf, valid_luhn
from lgpd_scanner.scanner import list_documents, scan_file


def names(text):
    return {p.name for p, _, _ in find_matches(text)}


class PatternTests(unittest.TestCase):
    def test_validators(self):
        self.assertTrue(valid_cpf("529.982.247-25"))
        self.assertFalse(valid_cpf("111.111.111-11"))
        self.assertFalse(valid_cpf("529.982.247-24"))
        self.assertTrue(valid_cnpj("11.222.333/0001-81"))
        self.assertTrue(valid_luhn("4111 1111 1111 1111"))

    def test_detection(self):
        t = ("CPF 529.982.247-25, email joao@exemplo.com.br, tel (11) 98765-4321, "
             "RG: 12.345.678-9, Data de nascimento: 05/03/1990, CEP 01310-100, "
             "cartão 4111 1111 1111 1111, CID F32.1, Rua das Flores, 123")
        found = names(t)
        for n in ("CPF", "E-mail", "Telefone", "RG", "Data de nascimento", "CEP",
                  "Cartão de crédito", "Saúde (CID / diagnóstico)", "Endereço residencial"):
            self.assertIn(n, found)

    def test_no_false_positive(self):
        self.assertNotIn("CPF", names("Protocolo 123.456.789-00 e 11111111111"))

    def test_mask(self):
        self.assertEqual(mask("529.982.247-25", "CPF"), "***.***.***-25")
        self.assertEqual(mask("joao@x.com", "E-mail"), "j***@x.com")


class ScanTests(unittest.TestCase):
    def test_files(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            doc = docx.Document()
            doc.add_paragraph("CPF: 529.982.247-25")
            doc.save(d / "a.docx")
            wb = openpyxl.Workbook()
            wb.active["B2"] = "maria@exemplo.com"
            wb.active["C3"] = 52998224725  # CPF numérico
            wb.save(d / "b.xlsx")
            (d / "c.pdf").write_bytes(b"not a pdf")
            self.assertEqual(len(list(list_documents([str(d)]))), 3)
            self.assertEqual(scan_file(d / "a.docx").findings[0].pattern, "CPF")
            self.assertEqual({f.pattern for f in scan_file(d / "b.xlsx").findings}, {"E-mail", "CPF"})
            self.assertIsNotNone(scan_file(d / "c.pdf").error)


if __name__ == "__main__":
    unittest.main()
