"""Extração de texto de arquivos DOCX, XLSX e PDF.

Cada extrator gera tuplas (localização, texto), permitindo indicar onde o dado foi achado.
"""
from __future__ import annotations

from pathlib import Path
from typing import Iterator

SUPPORTED = {".docx", ".xlsx", ".pdf"}


def extract_docx(path: Path) -> Iterator[tuple[str, str]]:
    import docx

    doc = docx.Document(str(path))
    for i, p in enumerate(doc.paragraphs, 1):
        if p.text.strip():
            yield f"parágrafo {i}", p.text
    for t, table in enumerate(doc.tables, 1):
        for r, row in enumerate(table.rows, 1):
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # células mescladas repetem
                    continue
                seen.add(id(cell._tc))
                if cell.text.strip():
                    yield f"tabela {t}, linha {r}", cell.text
    for s, section in enumerate(doc.sections, 1):
        for part, label in ((section.header, "cabeçalho"), (section.footer, "rodapé")):
            for p in part.paragraphs:
                if p.text.strip():
                    yield f"{label} (seção {s})", p.text


def extract_xlsx(path: Path) -> Iterator[tuple[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for r, row in enumerate(ws.iter_rows(values_only=True), 1):
                for c, value in enumerate(row, 1):
                    if value is None or value == "":
                        continue
                    # inteiros/floats grandes (CPF digitado como número) viram texto sem notação científica
                    if isinstance(value, float) and value.is_integer():
                        value = int(value)
                    yield f"planilha '{ws.title}', linha {r}, coluna {c}", str(value)
    finally:
        wb.close()


def extract_pdf(path: Path) -> Iterator[tuple[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ValueError("PDF protegido por senha")
        except Exception as exc:
            raise ValueError("PDF protegido por senha") from exc
    for n, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            yield f"página {n}", text


EXTRACTORS = {".docx": extract_docx, ".xlsx": extract_xlsx, ".pdf": extract_pdf}


def extract(path: Path) -> Iterator[tuple[str, str]]:
    return EXTRACTORS[path.suffix.lower()](path)
